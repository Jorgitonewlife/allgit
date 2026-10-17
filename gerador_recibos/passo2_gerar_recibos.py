# Bloco 2: gerar recibo de texto para cada cliente

import os

from openpyxl import load_workbook

caminho_planilha = os.path.join(os.path.dirname(__file__), "clientes.xlsx")

wb = load_workbook(caminho_planilha)
ws = wb.active

# Pula a primeira linha (cabeçalho)
linhas = list(ws.iter_rows(values_only=True))[1:]

for linha in linhas:
    nome, email, valor, servico = linha

    recibo = f"""==============================
        RECIBO DE PAGAMENTO
==============================

Cliente: {nome}
E-mail:  {email}
Serviço: {servico}
Valor:   R$ {valor}

Data: 15/10/2026

Assinatura: __________________________

==============================
Obrigado pela preferência!
=============================="""

    print(recibo)
    print()
