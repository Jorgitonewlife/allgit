# Bloco 1: ler planilha existente e mostrar o conteúdo

import os

from openpyxl import load_workbook

# Caminho da planilha de exemplo
caminho_planilha = os.path.join(os.path.dirname(__file__), "clientes.xlsx")

# Abre a planilha
wb = load_workbook(caminho_planilha)
ws = wb.active

print(f"Planilha: {caminho_planilha}")
print(f"Abas disponíveis: {wb.sheetnames}")
print(f"Total de linhas com dados: {ws.max_row}")
print(f"Total de colunas: {ws.max_column}")
print()

# Lê e mostra todas as linhas
for linha in ws.iter_rows(values_only=True):
    print(linha)
